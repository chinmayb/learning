"""Lab 3a — Parse docx/pdf -> data/parsed/docs.jsonl
Run: uv run python src/lab3_parse.py
One record per PDF page / per DOCX section:
  {doc_id, source, page, heading, text}
"""
from collections import Counter
from pathlib import Path

import docx
import fitz  # pymupdf
from common import DATA, write_jsonl


def parse_pdf(path: Path):
    doc = fitz.open(path)
    recs = []
    for i, page in enumerate(doc, start=1):
        text = page.get_text()  # try also: page.get_text("blocks") — compare!
        recs.append({"doc_id": path.stem, "source": path.name, "page": i,
                     "heading": "", "text": text})
    # TODO: fill 'heading' — hint: doc.get_toc() gives [level, title, page]
    return recs


def parse_docx(path: Path):
    d = docx.Document(path)
    recs, heading, buf = [], "", []

    def flush():
        if buf:
            recs.append({"doc_id": path.stem, "source": path.name, "page": None,
                         "heading": heading, "text": "\n".join(buf)})

    for para in d.paragraphs:
        if para.style.name.startswith("Heading"):
            flush()
            heading, buf = para.text.strip(), []
        elif para.text.strip():
            buf.append(para.text)
    flush()
    # TODO: d.tables are NOT in d.paragraphs. Your docs likely have parameter
    #       tables. Decide how to turn rows into text and add them.
    #       (row cells: [c.text for c in row.cells])
    return recs


def parse_markdown(path: Path):
    """Markdown / RST / text: split on heading lines. Clean structure for free."""
    recs, heading, buf = [], "", []
    for line in path.read_text(errors="ignore").splitlines():
        # TODO: detect headings — markdown '#', '##' ... ; RST uses underline rows (=====, -----)
        is_heading = line.startswith("#")
        if is_heading:
            if buf:
                recs.append({"doc_id": path.stem, "source": path.name, "page": None,
                             "heading": heading, "text": "\n".join(buf)})
            heading, buf = line.lstrip("# ").strip(), []
        else:
            buf.append(line)
    if buf:
        recs.append({"doc_id": path.stem, "source": path.name, "page": None,
                     "heading": heading, "text": "\n".join(buf)})
    # TODO: markdown image links / HTML tags / badges add noise. Strip them?
    return recs


def find_repeated_lines(recs, min_ratio=0.5):
    """Headers/footers show up on most pages. Return lines appearing on > min_ratio of pages."""
    # TODO: per doc_id, count each stripped line once per page; return the set above threshold
    ...


def clean(text: str, junk: set) -> str:
    # TODO: drop junk lines, fix 'resolu-\ntion' hyphenation, collapse blank lines
    # Careful: don't destroy CLI blocks / config snippets (indentation matters?)
    return text


if __name__ == "__main__":
    recs = []
    raw = DATA / "raw"
    for p in sorted(raw.rglob("*")):
        ext = p.suffix.lower()
        if ext == ".pdf":
            new = parse_pdf(p)
        elif ext == ".docx":
            new = parse_docx(p)
        elif ext in {".md", ".rst", ".txt"}:
            new = parse_markdown(p)
        else:
            continue   # HTML: save page as PDF, or convert to .md first (e.g. pandoc)
        # data/raw/<domain>/file -> domain; files directly in raw/ -> "misc"
        domain = p.relative_to(raw).parts[0] if len(p.relative_to(raw).parts) > 1 else "misc"
        for r in new:
            r["domain"] = domain
        recs += new
    junk = find_repeated_lines(recs) or set()
    print("junk lines found:", list(junk)[:20])
    for r in recs:
        r["text"] = clean(r["text"], junk)
    recs = [r for r in recs if r["text"].strip()]
    write_jsonl(DATA / "parsed" / "docs.jsonl", recs)
    print(f"{len(recs)} records from {len({r['doc_id'] for r in recs})} docs")
    # NOW: open 5 records next to the originals and list every defect in journal.md
